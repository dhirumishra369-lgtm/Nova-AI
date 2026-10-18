import datetime
import io
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd
import streamlit as st

st.set_page_config(
    page_title="Recipe Costing & Yield Calculator", page_icon="🍰", layout="wide"
)

# ----------------- SESSION STATE SETUP -----------------
if "ingredients" not in st.session_state:
  st.session_state.ingredients = pd.DataFrame([
      {"Ingredient": "Kaju (Cashew)", "Quantity_KG": 5.0, "Rate_Per_KG": 680.0},
      {"Ingredient": "Sugar", "Quantity_KG": 8.0, "Rate_Per_KG": 42.0},
      {"Ingredient": "Silver Vark", "Quantity_KG": 0.05, "Rate_Per_KG": 4000.0},
      {"Ingredient": "Cardamom / Ghee", "Quantity_KG": 0.2, "Rate_Per_KG": 600.0},
  ])

if "recipe_title" not in st.session_state:
  st.session_state.recipe_title = "Premium Kaju Katli"

if "master_excel_data" not in st.session_state:
  st.session_state.master_excel_data = None

if "price_lookup_dict" not in st.session_state:
  st.session_state.price_lookup_dict = {}

# Standard Nutritional Database per 100g for common ingredients
NUTRIENTS_DB = {
    "kaju": {"kcal": 553, "protein": 18.2, "fat": 43.8, "carbs": 30.2},
    "cashew": {"kcal": 553, "protein": 18.2, "fat": 43.8, "carbs": 30.2},
    "sugar": {"kcal": 387, "protein": 0.0, "fat": 0.0, "carbs": 100.0},
    "ghee": {"kcal": 900, "protein": 0.0, "fat": 100.0, "carbs": 0.0},
    "butter": {"kcal": 717, "protein": 0.85, "fat": 81.0, "carbs": 0.06},
    "milk": {"kcal": 42, "protein": 3.4, "fat": 1.0, "carbs": 5.0},
    "pistachio": {"kcal": 562, "protein": 20.0, "fat": 45.0, "carbs": 28.0},
    "almond": {"kcal": 579, "protein": 21.1, "fat": 49.9, "carbs": 21.6},
    "badam": {"kcal": 579, "protein": 21.1, "fat": 49.9, "carbs": 21.6},
    "cardamom": {"kcal": 311, "protein": 11.0, "fat": 7.0, "carbs": 68.0},
    "silver vark": {"kcal": 0.0, "protein": 0.0, "fat": 0.0, "carbs": 0.0},
    "salt": {"kcal": 0.0, "protein": 0.0, "fat": 0.0, "carbs": 0.0},
    "vanilla essence": {"kcal": 288, "protein": 0.1, "fat": 0.06, "carbs": 12.65},
}


def get_nutrients(ingredient_name):
  name_lower = str(ingredient_name).lower()
  for key, val in NUTRIENTS_DB.items():
    if key in name_lower:
      return val
  return {"kcal": 400, "protein": 5.0, "fat": 10.0, "carbs": 70.0}


# ----------------- PROFESSIONAL EXCEL GENERATORS -----------------
def generate_recipe_excel(
    recipe_name,
    df,
    loss_pct,
    final_yield,
    raw_cost,
    labor_cost,
    pack_cost,
    total_cost,
    cost_kg,
    margin_pct,
    sp_kg,
    profit_kg,
):
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = "Costing & Yield Report"
  ws.views.sheetView[0].showGridLines = True

  ws.merge_cells("A1:E1")
  title = ws["A1"]
  title.value = "RECIPE COSTING & BATCH YIELD REPORT"
  title.font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
  title.fill = PatternFill(
      start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
  )
  title.alignment = Alignment(horizontal="center", vertical="center")
  ws.row_dimensions[1].height = 35

  ws["A2"] = "Product / Recipe:"
  ws["B2"] = recipe_name
  ws["D2"] = "Date:"
  ws["E2"] = datetime.date.today().strftime("%d-%b-%Y")
  ws["A2"].font = Font(name="Calibri", size=11, bold=True, color="4B5563")
  ws["B2"].font = Font(name="Calibri", size=11, bold=True, color="111827")
  ws["D2"].font = Font(name="Calibri", size=11, bold=True, color="4B5563")
  ws["E2"].font = Font(name="Calibri", size=11, bold=True, color="111827")
  ws.row_dimensions[2].height = 22

  ws.merge_cells("A4:E4")
  sec1 = ws["A4"]
  sec1.value = "1. RAW MATERIAL & INGREDIENT BREAKDOWN"
  sec1.font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
  sec1.fill = PatternFill(
      start_color="DBEAFE", end_color="DBEAFE", fill_type="solid"
  )
  ws.row_dimensions[4].height = 24

  headers = [
      "S.No.",
      "Ingredient Name",
      "Quantity (KG)",
      "Rate / KG (₹)",
      "Total Amount (₹)",
  ]
  for col_idx, h in enumerate(headers, 1):
    c = ws.cell(row=5, column=col_idx, value=h)
    c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    c.fill = PatternFill(
        start_color="2563EB", end_color="2563EB", fill_type="solid"
    )
    c.alignment = Alignment(
        horizontal="center" if col_idx != 2 else "left", vertical="center"
    )
  ws.row_dimensions[5].height = 24

  thin_border = Border(
      left=Side(style="thin", color="E5E7EB"),
      right=Side(style="thin", color="E5E7EB"),
      top=Side(style="thin", color="E5E7EB"),
      bottom=Side(style="thin", color="E5E7EB"),
  )

  start_row = 6
  for idx, row in df.iterrows():
    curr = start_row + idx
    ws.row_dimensions[curr].height = 20
    ws.cell(row=curr, column=1, value=idx + 1).alignment = Alignment(
        horizontal="center"
    )
    ws.cell(row=curr, column=2, value=str(row["Ingredient"])).alignment = (
        Alignment(horizontal="left")
    )

    c3 = ws.cell(row=curr, column=3, value=float(row["Quantity_KG"]))
    c3.number_format = "#,##0.00"
    c3.alignment = Alignment(horizontal="right")

    c4 = ws.cell(row=curr, column=4, value=float(row["Rate_Per_KG"]))
    c4.number_format = "₹#,##0.00"
    c4.alignment = Alignment(horizontal="right")

    c5 = ws.cell(row=curr, column=5, value=f"=C{curr}*D{curr}")
    c5.number_format = "₹#,##0.00"
    c5.alignment = Alignment(horizontal="right")

    bg_color = "F9FAFB" if idx % 2 == 1 else "FFFFFF"
    for c_idx in range(1, 6):
      ws.cell(row=curr, column=c_idx).border = thin_border
      ws.cell(row=curr, column=c_idx).fill = PatternFill(
          start_color=bg_color, end_color=bg_color, fill_type="solid"
      )

  tot_r = start_row + len(df)
  ws.row_dimensions[tot_r].height = 22
  ws.cell(row=tot_r, column=2, value="Subtotal (Raw Materials)").font = Font(
      name="Calibri", size=11, bold=True
  )
  ws.cell(
      row=tot_r, column=3, value=f"=SUM(C{start_row}:C{tot_r-1})"
  ).font = Font(name="Calibri", size=11, bold=True)
  ws.cell(row=tot_r, column=3).number_format = "#,##0.00"
  ws.cell(row=tot_r, column=3).alignment = Alignment(horizontal="right")

  sub_amt = ws.cell(row=tot_r, column=5, value=f"=SUM(E{start_row}:E{tot_r-1})")
  sub_amt.font = Font(name="Calibri", size=11, bold=True)
  sub_amt.number_format = "₹#,##0.00"
  sub_amt.alignment = Alignment(horizontal="right")

  sum_border = Border(
      top=Side(style="thin", color="1E3A8A"),
      bottom=Side(style="double", color="1E3A8A"),
  )
  for c_idx in range(1, 6):
    ws.cell(row=tot_r, column=c_idx).border = sum_border
    ws.cell(row=tot_r, column=c_idx).fill = PatternFill(
        start_color="EFF6FF", end_color="EFF6FF", fill_type="solid"
    )

  sum_start = tot_r + 2
  ws.merge_cells(f"A{sum_start}:E{sum_start}")
  sec2 = ws[f"A{sum_start}"]
  sec2.value = "2. YIELD LOSS, OVERHEADS & FINANCIAL SUMMARY"
  sec2.font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
  sec2.fill = PatternFill(
      start_color="DBEAFE", end_color="DBEAFE", fill_type="solid"
  )
  ws.row_dimensions[sum_start].height = 24

  metrics = [
      ("Total Raw Material Batch Weight", f"=C{tot_r}", "KG", "#,##0.00"),
      ("Process / Moisture Loss (%)", loss_pct / 100.0, "%", "0.0%"),
      (
          "Final Net Yield Output",
          f"=E{sum_start+1}*(1-E{sum_start+2})",
          "KG",
          "#,##0.00",
      ),
      ("Raw Material Total Cost", f"=E{tot_r}", "INR", "₹#,##0.00"),
      ("Labour & Fuel Overheads", float(labor_cost), "INR", "₹#,##0.00"),
      ("Packaging & Box Cost", float(pack_cost), "INR", "₹#,##0.00"),
      (
          "Total Batch Production Cost",
          f"=SUM(E{sum_start+4}:E{sum_start+6})",
          "INR",
          "₹#,##0.00",
      ),
      (
          "Final Cost Per KG (True Cost)",
          f"=E{sum_start+7}/E{sum_start+3}",
          "INR/KG",
          "₹#,##0.00",
      ),
      ("Target Profit Margin (%)", margin_pct / 100.0, "%", "0.0%"),
      (
          "Suggested Selling Price Per KG",
          f"=E{sum_start+8}/(1-E{sum_start+9})",
          "INR/KG",
          "₹#,##0.00",
      ),
      (
          "Net Profit Per KG",
          f"=E{sum_start+10}-E{sum_start+8}",
          "INR/KG",
          "₹#,##0.00",
      ),
  ]

  for idx, (label, val, unit, num_fmt) in enumerate(metrics):
    r = sum_start + 1 + idx
    ws.row_dimensions[r].height = 21
    ws.merge_cells(f"A{r
